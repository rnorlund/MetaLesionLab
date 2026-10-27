"""Turn a picture-naming test that comes as a PowerPoint into a ZIP for the Naming Assessment.

Use it for a test you are licensed to use, such as the Philadelphia Naming Test obtained from its
publisher. The ZIP holds the pictures and a test.csv (file, target, alternates, practice); import
it on the Assessment page with "Import a test". It stays a private set in that browser: it is
never uploaded, listed or shared. Keep the ZIP private too, and never commit it: this script
refuses to write it inside the repository.

    python scripts/import_pptx_test.py test.pptx targets.txt -o my_test.zip [--practice 2] [--skip-slides 1]

Slides are read in order. On each slide the largest picture is the item; slides without a picture
(titles, instructions) are skipped. The target list has one line per picture slide, in the same
order: the target, then optionally a colon and accepted alternates separated by ";", e.g.

    cat: kitty; puss
    umbrella

It can also be a CSV with a header row: target, alternates, practice. --practice N marks the
first N items as practice items (given first, not scored).
"""
from __future__ import annotations

import argparse
import csv
import io
import sys
import zipfile
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]


def pictures_by_slide(pptx_path: Path, skip: set[int]) -> list[tuple[int, bytes, str]]:
    """(slide number, image bytes, extension) of the largest picture on each slide that has one."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
                continue
            try:
                image = sh.image  # pictures and filled picture placeholders
            except (AttributeError, ValueError, KeyError):  # other shapes, or an empty placeholder
                continue
            yield sh, image

    out = []
    for number, slide in enumerate(Presentation(str(pptx_path)).slides, 1):
        if number in skip:
            continue
        found = [(sh.width * sh.height, img) for sh, img in walk(slide.shapes)]
        if found:
            _, image = max(found, key=lambda t: t[0])
            out.append((number, image.blob, image.ext))
    return out


def read_targets(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8-sig")
    first = text.splitlines()[0].lower() if text.strip() else ""
    rows = []
    if path.suffix.lower() == ".csv" and "target" in first:
        for r in csv.DictReader(io.StringIO(text)):
            r = {k.strip().lower(): (v or "").strip() for k, v in r.items() if k}
            if r.get("target"):
                rows.append({"target": r["target"].lower(), "alternates": r.get("alternates", ""),
                             "practice": r.get("practice", "").lower() in ("1", "y", "yes", "true", "x", "practice")})
        return rows
    for line in text.splitlines():
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        target, _, alts = line.partition(":") if ":" in line else line.partition("\t")
        rows.append({"target": target.strip().lower(), "alternates": alts.strip(), "practice": False})
    return rows


def build_zip(pptx_path: Path, targets_path: Path, practice: int = 0, skip: set[int] | None = None) -> tuple[bytes, list[str]]:
    pictures = pictures_by_slide(pptx_path, skip or set())
    targets = read_targets(targets_path)
    if not pictures:
        raise SystemExit("No pictures found in the presentation.")
    if len(pictures) != len(targets):
        slides = ", ".join(str(n) for n, _, _ in pictures)
        raise SystemExit(f"{len(pictures)} slides have a picture (slides {slides}) but the target list has "
                         f"{len(targets)} entries. Make them match, or skip slides with --skip-slides.")
    log, buf = [], io.BytesIO()
    width = max(3, len(str(len(pictures))))
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        rows = []
        for k, ((slide, blob, ext), t) in enumerate(zip(pictures, targets), 1):
            name = f"{k:0{width}d}.{ext}"
            z.writestr(name, blob)
            is_practice = t["practice"] or k <= practice
            alts = "; ".join(a.strip().lower() for a in t["alternates"].replace("|", ";").split(";") if a.strip())
            rows.append({"file": name, "target": t["target"], "alternates": alts, "practice": "yes" if is_practice else ""})
            log.append(f"slide {slide:>3} -> {name}  {t['target']}" + (" (practice)" if is_practice else ""))
        out = io.StringIO()
        writer = csv.DictWriter(out, ["file", "target", "alternates", "practice"])
        writer.writeheader()
        writer.writerows(rows)
        z.writestr("test.csv", out.getvalue())
        z.writestr("README.txt", (
            f"Made by scripts/import_pptx_test.py from {pptx_path.name}.\n"
            "Import it on the Naming Assessment page (Import a test). It is kept as a private set in that browser.\n"
            "These pictures belong to the test's publisher: use them only as your license allows, keep this file\n"
            "private, and don't commit it to any repository.\n"))
    return buf.getvalue(), log


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("pptx", type=Path)
    ap.add_argument("targets", type=Path, help="target list: one line per picture slide, or a CSV")
    ap.add_argument("-o", "--out", type=Path, help="ZIP to write (default: <pptx name>_test.zip in the current folder)")
    ap.add_argument("--practice", type=int, default=0, help="mark the first N items as practice items")
    ap.add_argument("--skip-slides", default="", help="slide numbers to ignore, e.g. 1,2")
    ap.add_argument("--force", action="store_true", help="write inside the repository anyway (not recommended)")
    args = ap.parse_args()
    out = (args.out or Path.cwd() / f"{args.pptx.stem}_test.zip").resolve()
    if out.is_relative_to(REPO_ROOT) and not args.force:
        raise SystemExit(f"Refusing to write {out} inside the repository: test pictures must never be committed. "
                         "Choose a folder outside it with -o.")
    skip = {int(s) for s in args.skip_slides.replace(" ", "").split(",") if s}
    data, log = build_zip(args.pptx, args.targets, practice=args.practice, skip=skip)
    out.write_bytes(data)
    print("\n".join(log))
    print(f"Wrote {out} ({len(log)} items). Import it on the Naming Assessment page; keep the file private.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
